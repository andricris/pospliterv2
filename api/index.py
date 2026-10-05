"""
PO Splitter API — Vercel Serverless Function (Flask + openpyxl)
Pure Python, no COM dependency. Works in-memory (no temp files).
"""

import io
import json
import os
import traceback

from flask import Flask, request, jsonify, Response, send_from_directory
from flask_cors import CORS
from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection, Side

# Static folder = ../dist (frontend build output)
DIST_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'dist')
app = Flask(__name__, static_folder=DIST_DIR, static_url_path='')
CORS(app)

# Default column constants
DEFAULT_SOURCE_COL = 'AS'
DEFAULT_PO1_COL = 'AR'
DEFAULT_HNA_COL = 'F'
DEFAULT_MB_COL = 'P'
DEFAULT_ENABLE_PO1 = False

COL_AT = 46


def copy_cell_format(src_cell, dst_cell):
    """Copy cell format manually — more reliable than copy() on serverless."""
    try:
        s = src_cell.font
        dst_cell.font = Font(
            name=s.name, size=s.size, bold=s.bold, italic=s.italic,
            underline=s.underline, strike=s.strike, color=s.color
        )
    except: pass
    try:
        s = src_cell.fill
        if s and s.fill_type:
            dst_cell.fill = PatternFill(
                fill_type=s.fill_type, fgColor=s.fgColor, bgColor=s.bgColor
            )
    except: pass
    try:
        s = src_cell.border
        dst_cell.border = Border(
            left=s.left, right=s.right, top=s.top, bottom=s.bottom,
            diagonal=s.diagonal, diagonalUp=s.diagonalUp, diagonalDown=s.diagonalDown
        )
    except: pass
    try:
        dst_cell.number_format = src_cell.number_format
    except: pass
    try:
        s = src_cell.alignment
        dst_cell.alignment = Alignment(
            horizontal=s.horizontal, vertical=s.vertical, wrapText=s.wrapText,
            shrinkToFit=s.shrinkToFit, indent=s.indent
        )
    except: pass
    try:
        dst_cell.protection = Protection(locked=src_cell.protection.locked)
    except: pass


def find_sheet_and_header(wb):
    pp_name = None
    for name in wb.sheetnames:
        if name.upper() == 'PP' or 'PP' in name.upper():
            pp_name = name
            break
    if not pp_name:
        raise ValueError('Sheet "PP" tidak ditemukan.')
    ws = wb[pp_name]
    max_row = ws.max_row or 1
    max_col = ws.max_column or 50
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and 'po sisa' in cell.value.lower():
                header_row = cell.row
                break
        if header_row:
            break
    if not header_row:
        raise ValueError('Header "po sisa" tidak ditemukan.')
    return ws, pp_name, header_row


def to_number(value):
    """Coerce a cell value to float, or None if it is not numeric.
    Handles cached Excel errors (#VALUE!, #REF!, #N/A...) as non-numeric."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value)
        except (ValueError, TypeError):
            return None
    return None


def is_product_row(row):
    """Detect a product data row without relying on the NO column.

    The NO column is often a formula (e.g. =A6+1) whose cached value can be an
    Excel error (#VALUE!/#REF!) when the sheet is read with data_only=True.
    A row counts as data when it carries a PART NO / SUPPLIER / DESCRIPTION,
    or a positive numeric / formula in column A.
    """
    def _filled(v):
        return v is not None and str(v).strip() != ''

    a = row[0].value if len(row) > 0 else None
    b = row[1].value if len(row) > 1 else None
    c = row[2].value if len(row) > 2 else None
    d = row[3].value if len(row) > 3 else None

    if _filled(c) or _filled(b) or _filled(d):
        return True
    num = to_number(a)
    if num is not None and num > 0:
        return True
    if isinstance(a, str) and a.startswith('='):
        return True
    return False


def find_data_range(ws, header_row):
    max_row = ws.max_row or 1
    max_col = ws.max_column or 3
    data_start = None
    data_end = None
    for row in ws.iter_rows(min_row=header_row + 1, max_row=max_row,
                            max_col=max(max_col, 4)):
        if is_product_row(row):
            row_num = row[0].row
            if data_start is None:
                data_start = row_num
            data_end = row_num
    if data_start is None:
        raise ValueError('Data produk tidak ditemukan.')
    return data_start, data_end


def process_excel(file_bytes, num_po, percentages, source_col='AS', po1_col='AR',
                  enable_po1=False, hna_col='F', mb_col='P'):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, column_index_from_string

    wb = load_workbook(io.BytesIO(file_bytes))
    ws, pp_name, header_row = find_sheet_and_header(wb)

    if ws.protection.sheet:
        ws.protection.sheet = False
    if ws.protection.password:
        ws.protection.password = None

    data_start, data_end = find_data_range(ws, header_row)

    src_idx = column_index_from_string(source_col)
    po1_idx = column_index_from_string(po1_col)
    hna_idx = column_index_from_string(hna_col)
    mb_idx = column_index_from_string(mb_col)

    start_col = COL_AT
    summary_row_val = header_row - 2
    total_col = start_col + num_po
    data_number_format = '#,##0'

    ref_hdr = ws.cell(row=header_row, column=1)
    for i in range(num_po + 1):
        copy_cell_format(ref_hdr, ws.cell(row=header_row, column=start_col + i))

    for i in range(num_po):
        ws.cell(row=header_row, column=start_col + i).value = f"PO Sisa {i + 1} ({percentages[i]}%)"
    ws.cell(row=header_row, column=total_col).value = "TOTAL PO Sisa"

    src_letter = get_column_letter(src_idx)
    po1_letter = get_column_letter(po1_idx)
    mb_letter = get_column_letter(mb_idx)
    hna_letter = get_column_letter(hna_idx)

    def source_expr(row_num):
        if enable_po1:
            return f"({src_letter}{row_num}-{po1_letter}{row_num})"
        return f"{src_letter}{row_num}"

    for row_num in range(data_start, data_end + 1):
        s_expr = source_expr(row_num)

        # Running sum of previous PO columns (as Excel refs), built incrementally.
        # Cap each PO with MIN(MROUND(...), source - prev_sum) so cumulative total
        # can never exceed the source qty -> guarantees no minus on the last PO.
        prev_refs = []
        for i in range(num_po):
            col_idx = start_col + i
            cell = ws.cell(row=row_num, column=col_idx)
            col_letter = get_column_letter(col_idx)
            pct = percentages[i] / 100

            if i == num_po - 1:
                # Last PO = source - SUM(previous POs). Never negative because
                # every previous PO was capped to source - prev_sum.
                if prev_refs:
                    cell.value = f"={s_expr}-({'+'.join(prev_refs)})"
                else:
                    # Single PO case: all qty goes here
                    cell.value = f"={s_expr}"
            else:
                # PO 1..N-1: MROUND to MB multiple, capped to remaining available qty.
                # remaining = source - SUM(prev_refs) -> guaranteed >= 0.
                mround_part = f"MROUND({s_expr}*{pct},{mb_letter}{row_num})"
                if prev_refs:
                    cell.value = f"=MIN({mround_part},{s_expr}-({'+'.join(prev_refs)}))"
                else:
                    cell.value = f"=MIN({mround_part},{s_expr})"
            cell.number_format = data_number_format
            prev_refs.append(f"{col_letter}{row_num}")

        total_cell = ws.cell(row=row_num, column=total_col)
        po_refs = [f"{get_column_letter(start_col + j)}{row_num}" for j in range(num_po)]
        total_cell.value = f"={'+'.join(po_refs)}"
        total_cell.number_format = data_number_format

    ws.cell(row=summary_row_val, column=1).value = "TOTAL NILAI RUPIAH (Qty x HNA)"
    for i in range(num_po):
        cl = get_column_letter(start_col + i)
        ws.cell(row=summary_row_val, column=start_col + i).value = \
            f"=SUMPRODUCT({cl}{data_start}:{cl}{data_end},{hna_letter}{data_start}:{hna_letter}{data_end})"
    tcl = get_column_letter(total_col)
    ws.cell(row=summary_row_val, column=total_col).value = \
        f"=SUMPRODUCT({tcl}{data_start}:{tcl}{data_end},{hna_letter}{data_start}:{hna_letter}{data_end})"

    output = io.BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return output.getvalue()


def preview_excel(file_bytes):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws, pp_name, header_row = find_sheet_and_header(wb)
    max_col = ws.max_column or 50

    headers = []
    for col_idx in range(1, max_col + 1):
        cell_val = ws.cell(row=header_row, column=col_idx).value
        if cell_val and str(cell_val).strip():
            letter = get_column_letter(col_idx)
            headers.append({
                'col': col_idx,
                'letter': letter,
                'label': f"{letter} — {str(cell_val).strip()}"
            })

    products = []
    seq = 0
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        if not is_product_row(row):
            continue

        # NO column may hold an Excel error (#VALUE!/#REF!) when the sheet is
        # read with data_only=True. Fall back to a sequential row counter.
        a_val = row[0].value
        num = to_number(a_val)
        seq += 1
        no = int(num) if (num is not None and num > 0) else seq

        vals = []
        for cell in row:
            v = cell.value
            if v is None:
                vals.append(0)
            elif isinstance(v, (int, float)):
                vals.append(v)
            else:
                n = to_number(v)
                vals.append(n if n is not None else 0)

        def str_val(v):
            return str(v) if v is not None else ''

        products.append({
            'no': no,
            'supplier': str_val(row[1].value) if len(row) > 1 else '',
            'partNo': str_val(row[2].value) if len(row) > 2 else '',
            'desc': str_val(row[3].value) if len(row) > 3 else '',
            'vals': vals,
        })

    wb.close()
    return products, pp_name, headers


# ============================================================
# API Routes (both /api/* and /* for Vercel routing)
# ============================================================

@app.route('/', methods=['GET'])
def index():
    """Serve frontend index.html from dist/ folder."""
    return send_from_directory(DIST_DIR, 'index.html')


@app.route('/<path:path>', methods=['GET'])
def static_files(path):
    """Serve static assets (JS, CSS, images) from dist/ folder.
    Falls back to index.html for unknown paths (SPA support)."""
    full = os.path.join(DIST_DIR, path)
    if os.path.isfile(full):
        return send_from_directory(DIST_DIR, path)
    return send_from_directory(DIST_DIR, 'index.html')


@app.route('/api/health', methods=['GET'])
@app.route('/health', methods=['GET'])
def health():
    return jsonify({'status': 'ok'})


@app.route('/api/preview', methods=['POST'])
@app.route('/preview', methods=['POST'])
def api_preview():
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file uploaded'}), 400
        file = request.files['file']
        products, sheet_name, headers = preview_excel(file.read())
        return jsonify({
            'products': products,
            'sheetName': sheet_name,
            'headers': headers,
            'defaults': {
                'sourceCol': DEFAULT_SOURCE_COL,
                'po1Col': DEFAULT_PO1_COL,
                'enablePO1': DEFAULT_ENABLE_PO1,
                'hnaCol': DEFAULT_HNA_COL,
                'mbCol': DEFAULT_MB_COL,
            }
        })
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@app.route('/api/process', methods=['POST'])
@app.route('/process', methods=['POST'])
def api_process():
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No file uploaded'}), 400

        file = request.files['file']
        num_po = int(request.form.get('numPO', 4))
        percentages = json.loads(str(request.form.get('percentages', '[]')))
        source_col = request.form.get('sourceCol', DEFAULT_SOURCE_COL).upper()
        po1_col = request.form.get('po1Col', DEFAULT_PO1_COL).upper()
        enable_po1 = request.form.get('enablePO1', 'false').lower() == 'true'
        hna_col = request.form.get('hnaCol', DEFAULT_HNA_COL).upper()
        mb_col = request.form.get('mbCol', DEFAULT_MB_COL).upper()

        if len(percentages) < num_po:
            return jsonify({'error': 'Persentase tidak lengkap'}), 400
        total_pct = sum(percentages[:num_po])
        if total_pct != 100:
            return jsonify({'error': f'Total persentase harus 100%, sekarang {total_pct}%'}), 400

        file_bytes = process_excel(file.read(), num_po, percentages,
                                   source_col, po1_col, enable_po1, hna_col, mb_col)

        return Response(
            file_bytes,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename=PP_Output_PO_Split_{num_po}PO.xlsx'}
        )
    except Exception as e:
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500
